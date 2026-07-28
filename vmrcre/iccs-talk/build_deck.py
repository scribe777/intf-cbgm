#!/usr/bin/env python3
"""Build the ICCS XIII talk deck on top of the team's Göttingen template.

Self-contained: every asset it needs lives next to this script.
  python3 build_deck.py   ->  writes ai-integration-slides.pptx here.

Requires: python-pptx  (pip install python-pptx)
Assets in this folder:
  template.pptx                 the team's title-slide template (logos + warm bg)
  logos/image[1-4].png          emblem, Uni Göttingen, NAWG, IACS  (closing slide)
  slide_*.png, tesla_*.png      demo screenshots (Ps 22:2 collation)
See CBGM_SCREENSHOTS_TODO.md for the one open item (real CBGM screenshots).
"""
import os, struct
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE  = os.path.dirname(os.path.abspath(__file__))
TPL   = os.path.join(BASE, "template.pptx")
OUT   = os.path.join(BASE, "ai-integration-slides.pptx")
SHOTS = BASE + "/"
LOGOS = os.path.join(BASE, "logos") + "/"

WARM = RGBColor(0xFF,0xF7,0xF6); INK = RGBColor(0x1A,0x1A,0x1A)
BLUE = RGBColor(0x1F,0x5C,0x99); GOLD = RGBColor(0xC9,0x9A,0x2E)
GRAY = RGBColor(0x8A,0x83,0x80); WHITE= RGBColor(0xFF,0xFF,0xFF)
TITLEFONT="Roboto"; BODYFONT="Calibri"

prs = Presentation(TPL)
blank = min(prs.slide_layouts, key=lambda L: len(L.placeholders))

def png_size(p):
    with open(p,'rb') as f: head=f.read(24)
    return struct.unpack('>II', head[16:24])

def set_run(r, text, size, color=INK, font=BODYFONT, bold=False, italic=False):
    r.text=text; f=r.font; f.size=Pt(size); f.name=font; f.bold=bold; f.italic=italic
    f.color.rgb=color

def warm_bg(slide):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=WARM

def add_title(slide, text, size=28):
    tb=slide.shapes.add_textbox(Inches(0.62), Inches(0.42), Inches(12.1), Inches(1.0))
    tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; set_run(p.add_run(), text, size, INK, TITLEFONT)
    rule=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.42), Inches(2.6), Pt(3))
    rule.fill.solid(); rule.fill.fore_color.rgb=BLUE; rule.line.fill.background(); rule.shadow.inherit=False
    return tb

def footer(slide, n):
    tb=slide.shapes.add_textbox(Inches(0.62), Inches(7.02), Inches(9.0), Inches(0.35))
    set_run(tb.text_frame.paragraphs[0].add_run(),
            "AI integration with collaborative digital tools for critical editions  ·  ICCS XIII, Göttingen 2026",
            9, GRAY, BODYFONT)
    tb2=slide.shapes.add_textbox(Inches(12.0), Inches(7.02), Inches(0.9), Inches(0.35))
    p2=tb2.text_frame.paragraphs[0]; p2.alignment=PP_ALIGN.RIGHT
    set_run(p2.add_run(), str(n), 9, GRAY, BODYFONT)

def content_slide(title, n, tsize=28):
    s=prs.slides.add_slide(blank)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    warm_bg(s); add_title(s, title, tsize); footer(s, n)
    return s

def add_image_fit(slide, path, box_l, box_t, box_w, box_h, caption=None):
    w,h=png_size(path); ar=w/h
    if box_w/box_h > ar: ih=box_h; iw=box_h*ar
    else: iw=box_w; ih=box_w/ar
    il=box_l+(box_w-iw)/2; it=box_t+(box_h-ih)/2
    pic=slide.shapes.add_picture(path, Inches(il), Inches(it), Inches(iw), Inches(ih))
    pic.line.color.rgb=RGBColor(0xDD,0xD5,0xD3); pic.line.width=Pt(0.75)
    if caption:
        tb=slide.shapes.add_textbox(Inches(box_l), Inches(box_t+box_h+0.02), Inches(box_w), Inches(0.4))
        p=tb.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        set_run(p.add_run(), caption, 13, GRAY, BODYFONT, italic=True)
    return pic

def bullets(slide, items, left=0.72, top=1.75, width=11.9, height=4.9, size=18, gap=8):
    tb=slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf=tb.text_frame; tf.word_wrap=True; first=True
    for head, rest in items:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        p.space_after=Pt(gap)
        set_run(p.add_run(), "•  ", size, BLUE, BODYFONT, bold=True)
        set_run(p.add_run(), head, size, INK, BODYFONT, bold=True)
        if rest: set_run(p.add_run(), " — "+rest, size, INK, BODYFONT)
    return tb

def big_center(slide, text, size=30, top=2.2, height=3.0, color=INK, font=TITLEFONT):
    tb=slide.shapes.add_textbox(Inches(1.1), Inches(top), Inches(11.1), Inches(height))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    set_run(p.add_run(), text, size, color, font)
    return tb

def pill(slide, text, top, color=BLUE, txtcolor=WHITE, size=18, width=9.0):
    l=Inches((13.333-width)/2)
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(top), Inches(width), Inches(0.7))
    sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background(); sh.shadow.inherit=False
    tf=sh.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    set_run(p.add_run(), text, size, txtcolor, TITLEFONT, bold=True)
    return sh

# ===================================================== SLIDE 1  (edit template title slide)
s1=prs.slides[0]
for sh in s1.shapes:
    if sh.has_text_frame and sh.name.startswith("Titel"):
        tf=sh.text_frame; tf.clear()
        set_run(tf.paragraphs[0].add_run(), "AI integration with collaborative", 32, INK, TITLEFONT)
        set_run(tf.add_paragraph().add_run(), "digital tools for critical editions", 32, INK, TITLEFONT)
        sh.top=Inches(1.95); sh.height=Inches(1.7)
sub=s1.shapes.add_textbox(Inches(0.71), Inches(3.7), Inches(11.9), Inches(1.1))
tf=sub.text_frame; tf.word_wrap=True
set_run(tf.paragraphs[0].add_run(), "Troy A. Griffitts", 22, BLUE, TITLEFONT, bold=True)
set_run(tf.add_paragraph().add_run(), "Virtual Manuscript Room Collaborative Research Environment", 16, INK, BODYFONT)
pc=tf.add_paragraph(); pc.space_before=Pt(6)
set_run(pc.add_run(), "13th International Congress of Coptic Studies  ·  Göttingen  ·  Friday 31 July 2026", 14, GRAY, BODYFONT)

# ===================================================== SLIDE 2  pipeline (CBGM = gold frontier)
s=content_slide("One edition, eleven stages", 2)
stages=[("Cataloguing",0),("Imaging",0),("Indexing",1),("Transcribing",1),
        ("Collating",1),("Regularizing",1),("Setting variant units",1),("Versional evidence",1),
        ("Ordering readings",0),("Textual history (CBGM)",2),("Publication",0)]
cols=4; cw=2.85; ch=0.82; gx=0.22; gy=0.30
gridw=cols*cw+(cols-1)*gx; startx=(13.333-gridw)/2; starty=1.9
for i,(name,hot) in enumerate(stages):
    x=startx+(i%cols)*(cw+gx); y=starty+(i//cols)*(ch+gy)
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(ch))
    sh.shadow.inherit=False
    if hot==1:   sh.fill.solid(); sh.fill.fore_color.rgb=BLUE; sh.line.fill.background(); tc=WHITE
    elif hot==2: sh.fill.solid(); sh.fill.fore_color.rgb=GOLD; sh.line.fill.background(); tc=INK
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb=WHITE
        sh.line.color.rgb=RGBColor(0xCF,0xC7,0xC5); sh.line.width=Pt(1); tc=GRAY
    p=sh.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    set_run(p.add_run(), f"{i+1}. {name}", 14, tc, BODYFONT, bold=bool(hot))
leg=s.shapes.add_textbox(Inches(startx), Inches(starty+3*(ch+gy)+0.15), Inches(gridw), Inches(0.5))
lp=leg.text_frame.paragraphs[0]; lp.alignment=PP_ALIGN.CENTER
set_run(lp.add_run(), "Blue = AI integrated today (6)", 14, BLUE, BODYFONT, bold=True)
set_run(lp.add_run(), "        Gold = CBGM local stemma — just landed", 14, GOLD, BODYFONT, bold=True)

# ===================================================== SLIDE 3  question
s=content_slide("The question", 3)
big_center(s, "Can AI be a legitimate assistant at every stage — without eroding the accountability that makes an edition authoritative?", 30, top=1.9, height=2.6)
pill(s, "Assist everywhere.  Author nothing that becomes the edition.", 5.1)

# ===================================================== SLIDE 4  six points
s=content_slide("Six integration points, today", 4)
bullets(s, [
 ("Indexing","locate the biblical passage on each manuscript page"),
 ("Transcribing","first-pass diplomatic text — useful even where scholars will (rightly) challenge it"),
 ("Collating","language-aware alignment across Hebrew · LXX Greek · Latin · Coptic · Syriac"),
 ("Regularizing","propose orthographic normalisations for human review"),
 ("Setting variant units","group aligned columns into variation units"),
 ("Versional evidence","align a translation back to its Greek / Hebrew Vorlage"),
], size=18, gap=10)

# ===================================================== SLIDE 5  demo intro
s=content_slide("Live demo — Psalm 22:2, five languages", 5)
add_image_fit(s, SHOTS+"tesla_ps22_2_landing.png", 2.9, 1.7, 7.5, 5.05,
              caption="Drew Longacre's Hebrew Bible Critical Edition of the Psalms, in the VMRCRE collation editor")

# ===================================================== SLIDE 6  before collatex
s=content_slide("Before — CollateX alignment", 6)
add_image_fit(s, SHOTS+"slide_01_collation_console.png", 2.9, 1.7, 7.5, 5.05,
              caption="Baseless, token-level: the alignment sprawls to 37 columns")

# ===================================================== SLIDE 7  after gemini
s=content_slide("After — language-aware AI alignment", 7)
add_image_fit(s, SHOTS+"slide_06_gemini_after.png", 2.9, 1.7, 7.5, 5.05,
              caption="Same verse, 26 columns · $0.30 · 2m39s · every run logged with model + cost")

# ===================================================== SLIDE 8  both tables
s=content_slide("The width is the argument", 8)
add_image_fit(s, SHOTS+"slide_09_both_tables.png", 2.9, 1.7, 7.5, 5.05,
              caption="37 → 26 columns: fewer, better-motivated variation units")

# ===================================================== SLIDE 9  prose
s=content_slide("It shows its reasoning", 9)
add_image_fit(s, SHOTS+"slide_07_ai_prose.png", 2.9, 1.7, 7.5, 5.05,
              caption="Groups ⲟ ⲑⲉⲟⲥ ⲙⲟⲩ to Hebrew אלי — and states a Vorlage judgment openly")

# ===================================================== SLIDE 10  gate the output
s=content_slide("Lesson 1 — gate every answer", 10)
big_center(s, "AI is eager to please, and guesses persuasively.\nNever trust the prose — check it against a mechanical invariant.",
           24, top=1.7, height=2.0)
pill(s, "Collation gate:  tokens missing  ·  tokens added  ·  input words silently changed", 4.9, color=GOLD, txtcolor=INK, size=16)
cap=s.shapes.add_textbox(Inches(1.1), Inches(5.75), Inches(11.1), Inches(0.7))
cp=cap.text_frame.paragraphs[0]; cp.alignment=PP_ALIGN.CENTER; cap.text_frame.word_wrap=True
set_run(cp.add_run(), "If any input token was altered without being reported, reject the whole result.", 15, GRAY, BODYFONT, italic=True)

# ===================================================== SLIDE 11  review gate
s=content_slide("AI proposes; a human accepts", 11)
add_image_fit(s, SHOTS+"slide_10_suggestions.png", 2.9, 1.7, 7.5, 5.05,
              caption="Regularisation suggestions — each Accept / Dismiss by an editor. Nothing auto-commits.")

# ===================================================== SLIDE 12  AI writes tools
s=content_slide("Lesson 2 — prefer AI that writes tools", 12)
big_center(s, "AI is non-deterministic. Move the non-determinism to one-time authoring:", 24, top=1.7, height=1.4)
bullets(s, [
 ("Don't","ask the model to redo the task by hand every time"),
 ("Do","have it write deterministic code you review, keep, and improve"),
], left=1.6, top=3.3, width=10.0, size=20, gap=14)

# ===================================================== SLIDE 13  AI as member
s=content_slide("AI joins the contribution ethos", 13)
bullets(s, [
 ("A named, accountable member","its own userID; every action in the usage log — who, what, which model, what it cost"),
 ("Bring-Your-Own-Key","three tiers, user → project → system; whose budget pays is always explicit"),
 ("Enriches the shared research dataset","many teams draw on it for their own projects — it is not poured into one edition"),
], size=19, gap=14)

# ===================================================== SLIDE 14  CBGM — the pattern generalizes
s=content_slide("The pattern generalizes — CBGM local stemma", 14, tsize=26)
steps=["Deterministic\ncoherence digest","Model registry\n· BYOK","AI proposes\nonly the DAG","Edge-by-edge review\ncost + tokens logged"]
pw=2.75; ph=1.05; gap=0.30; n=len(steps)
total=n*pw+(n-1)*gap; sx=(13.333-total)/2; y=1.95
for i,txt in enumerate(steps):
    x=sx+i*(pw+gap)
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(pw), Inches(ph))
    sh.shadow.inherit=False; sh.fill.solid(); sh.fill.fore_color.rgb=BLUE; sh.line.fill.background()
    tf=sh.text_frame; tf.word_wrap=True
    for j,line in enumerate(txt.split("\n")):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.alignment=PP_ALIGN.CENTER
        set_run(p.add_run(), line, 12.5, WHITE, BODYFONT, bold=(j==0))
    if i<n-1:
        ar=s.shapes.add_textbox(Inches(x+pw-0.02), Inches(y+ph/2-0.25), Inches(gap+0.06), Inches(0.5))
        pa=ar.text_frame.paragraphs[0]; pa.alignment=PP_ALIGN.CENTER
        set_run(pa.add_run(), "→", 20, GRAY, BODYFONT, bold=True)
bullets(s, [
 ("The hardest judgment in the pipeline","for each variation unit, which reading is the source of which — the CBGM local stemma (locstem)"),
 ("Grounded on evidence it cannot invent","seeded from pre-genealogical coherence; the model proposes only the edges (Lesson 1)"),
 ("A first-class NTVMR citizen","logs in with NTVMR credentials; decisions round-trip to the shared project-data store"),
], left=0.9, top=3.5, width=11.5, size=17, gap=10)
tk=s.shapes.add_textbox(Inches(0.9), Inches(6.15), Inches(11.5), Inches(0.5))
tp=tk.text_frame.paragraphs[0]; tp.alignment=PP_ALIGN.CENTER
set_run(tp.add_run(), "Same architecture as collation — a new stage, not a new system.", 15, BLUE, TITLEFONT, bold=True)

# ===================================================== SLIDE 15  closing
s=prs.slides.add_slide(blank)
for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
warm_bg(s)
big_center(s, "Assist everywhere.  Author nothing.", 40, top=1.7, height=1.6)
big_center(s, "AI as a legitimate assistant across the editing pipeline.", 22, top=3.15, height=1.0, color=BLUE)
ct=s.shapes.add_textbox(Inches(1.1), Inches(4.2), Inches(11.1), Inches(0.9)); ct.text_frame.word_wrap=True
p=ct.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
set_run(p.add_run(), "Troy A. Griffitts  ·  VMRCRE / NTVMR  ·  coptot.manuscriptroom.com", 16, GRAY, BODYFONT)
s.shapes.add_picture(LOGOS+"image2.png", Inches(0.8), Inches(6.0), height=Inches(0.62))
s.shapes.add_picture(LOGOS+"image4.png", Inches(5.7), Inches(5.75), height=Inches(1.0))
s.shapes.add_picture(LOGOS+"image3.png", Inches(9.1), Inches(5.9), height=Inches(0.85))

prs.save(OUT)
print("saved", OUT, "— slides:", len(prs.slides._sldIdLst))
